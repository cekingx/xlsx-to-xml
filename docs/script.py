import openpyxl
import sys
from xml.sax.saxutils import escape

def cell(v):
    """Convert a cell value to string, treating None as empty string."""
    if v is None:
        return ""
    if isinstance(v, float):
        # Preserve integer-looking floats without a spurious decimal,
        # otherwise round to 2 decimals like DJP's own converter output.
        if v == int(v):
            return str(int(v))
        return f"{v:.2f}"
    return str(v)

def convert(xlsx_path, xml_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws_f = wb["Faktur"]
    ws_d = wb["DetailFaktur"]

    # TIN: NPWP Penjual, from row 1, column C (3rd col)
    tin = cell(ws_f.cell(row=1, column=3).value)

    # Read Faktur rows (header at row 3, data starts row 4)
    faktur_rows = []
    for row in ws_f.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        faktur_rows.append(row)

    # Read DetailFaktur rows (header at row 1, data starts row 2), group by Baris
    detail_by_baris = {}
    for row in ws_d.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        baris = row[0]
        detail_by_baris.setdefault(baris, []).append(row)

    lines = []
    lines.append('<?xml version="1.0" encoding="utf-8"?>')
    lines.append('<TaxInvoiceBulk xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:noNamespaceSchemaLocation="TaxInvoice.xsd">')
    lines.append(f'  <TIN>{escape(tin)}</TIN>')
    lines.append('  <ListOfTaxInvoice>')

    for row in faktur_rows:
        (baris, tgl_faktur, jenis_faktur, kode_trx, ket_tambahan, dok_pendukung,
         referensi, cap_fasilitas, id_tku_penjual, npwp_pembeli, jenis_id_pembeli,
         negara_pembeli, no_dok_pembeli, nama_pembeli, alamat_pembeli, email_pembeli,
         id_tku_pembeli) = row[:17]

        # Date formatting: take date part only
        if hasattr(tgl_faktur, 'strftime'):
            tgl_str = tgl_faktur.strftime('%Y-%m-%d')
        else:
            tgl_str = cell(tgl_faktur)[:10]

        lines.append('    <TaxInvoice>')
        lines.append(f'      <TaxInvoiceDate>{escape(tgl_str)}</TaxInvoiceDate>')
        lines.append(f'      <TaxInvoiceOpt>{escape(cell(jenis_faktur))}</TaxInvoiceOpt>')
        lines.append(f'      <TrxCode>{escape(cell(kode_trx))}</TrxCode>')
        lines.append(f'      <AddInfo>{escape(cell(ket_tambahan))}</AddInfo>')
        lines.append(f'      <CustomDoc>{escape(cell(dok_pendukung))}</CustomDoc>')
        lines.append(f'      <RefDesc>{escape(cell(referensi))}</RefDesc>')
        lines.append(f'      <FacilityStamp>{escape(cell(cap_fasilitas))}</FacilityStamp>')
        lines.append(f'      <SellerIDTKU>{escape(cell(id_tku_penjual))}</SellerIDTKU>')
        lines.append(f'      <BuyerTin>{escape(cell(npwp_pembeli))}</BuyerTin>')
        lines.append(f'      <BuyerDocument>{escape(cell(jenis_id_pembeli))}</BuyerDocument>')
        lines.append(f'      <BuyerCountry>{escape(cell(negara_pembeli))}</BuyerCountry>')
        lines.append(f'      <BuyerDocumentNumber>{escape(cell(no_dok_pembeli))}</BuyerDocumentNumber>')
        lines.append(f'      <BuyerName>{escape(cell(nama_pembeli))}</BuyerName>')
        lines.append(f'      <BuyerAdress>{escape(cell(alamat_pembeli))}</BuyerAdress>')
        lines.append(f'      <BuyerEmail>{escape(cell(email_pembeli))}</BuyerEmail>')
        lines.append(f'      <BuyerIDTKU>{escape(cell(id_tku_pembeli))}</BuyerIDTKU>')
        lines.append('      <ListOfGoodService>')

        for drow in detail_by_baris.get(baris, []):
            (d_baris, opt, kode_barang, nama_barang, satuan, harga, jumlah,
             diskon, dpp, dpp_lain, tarif_ppn, ppn, tarif_ppnbm, ppnbm) = drow[:14]
            lines.append('        <GoodService>')
            lines.append(f'          <Opt>{escape(cell(opt))}</Opt>')
            lines.append(f'          <Code>{escape(cell(kode_barang))}</Code>')
            lines.append(f'          <Name>{escape(cell(nama_barang))}</Name>')
            lines.append(f'          <Unit>{escape(cell(satuan))}</Unit>')
            lines.append(f'          <Price>{escape(cell(harga))}</Price>')
            lines.append(f'          <Qty>{escape(cell(jumlah))}</Qty>')
            lines.append(f'          <TotalDiscount>{escape(cell(diskon))}</TotalDiscount>')
            lines.append(f'          <TaxBase>{escape(cell(dpp))}</TaxBase>')
            lines.append(f'          <OtherTaxBase>{escape(cell(dpp_lain))}</OtherTaxBase>')
            lines.append(f'          <VATRate>{escape(cell(tarif_ppn))}</VATRate>')
            lines.append(f'          <VAT>{escape(cell(ppn))}</VAT>')
            lines.append(f'          <STLGRate>{escape(cell(tarif_ppnbm))}</STLGRate>')
            lines.append(f'          <STLG>{escape(cell(ppnbm))}</STLG>')
            lines.append('        </GoodService>')

        lines.append('      </ListOfGoodService>')
        lines.append('    </TaxInvoice>')

    lines.append('  </ListOfTaxInvoice>')
    lines.append('</TaxInvoiceBulk>')

    with open(xml_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')

if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
